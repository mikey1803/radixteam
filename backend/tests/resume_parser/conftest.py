"""Shared fixtures for Resume Parser tests."""

from __future__ import annotations

import io
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document

from backend.modules.resume_parser.models import ParsedResumeRecord
from backend.modules.resume_parser.repository import ResumeRepository
from backend.modules.resume_parser.schemas import ParseMetadata, ParseResult, ParseStatus, ParsedResume, FileType
from backend.modules.resume_parser.service import ResumeParserService


# ── Sample Data ─────────────────────────────────────────────────────


@pytest.fixture
def sample_parsed_resume() -> ParsedResume:
    return ParsedResume(
        personal_info={
            "first_name": "Jane",
            "last_name": "Doe",
            "email": "jane@example.com",
            "phone": "+15551234567",
            "location": "San Francisco, CA",
            "headline": "Senior Software Engineer",
        },
        summary="Experienced software engineer with 5+ years of experience.",
        education=[
            {
                "institution": "MIT",
                "degree": "BS",
                "field_of_study": "Computer Science",
                "start_date": "2015-09",
                "end_date": "2019-06",
                "gpa": 3.8,
            }
        ],
        experience=[
            {
                "company": "Acme Corp",
                "title": "Senior Engineer",
                "location": "San Francisco, CA",
                "start_date": "2021-01",
                "end_date": None,
                "description": "Leading backend development.",
                "highlights": ["Built microservices", "Reduced latency 40%"],
            }
        ],
        skills=[
            {"name": "Python", "category": "Programming", "proficiency": "Expert"},
            {"name": "FastAPI", "category": "Framework", "proficiency": "Advanced"},
        ],
        projects=[
            {
                "name": "OpenAnalytics",
                "description": "Open-source analytics dashboard",
                "url": "https://github.com/janedoe/openanalytics",
                "technologies": ["Python", "React"],
                "highlights": ["500+ GitHub stars"],
            }
        ],
        certifications=[
            {
                "name": "AWS Solutions Architect",
                "issuer": "AWS",
                "date_obtained": "2023-03",
                "expiry_date": None,
                "url": None,
            }
        ],
        links=[
            {"label": "LinkedIn", "url": "https://linkedin.com/in/janedoe"},
            {"label": "GitHub", "url": "https://github.com/janedoe"},
        ],
    )


@pytest.fixture
def sample_metadata() -> ParseMetadata:
    return ParseMetadata(
        parse_id="test-parse-001",
        filename="resume.pdf",
        file_type=FileType.PDF,
        file_size_bytes=1024,
        raw_text_length=500,
        ai_model="gpt-4o-mini",
        parser_version="1.0.0",
        parse_duration_ms=150.5,
    )


@pytest.fixture
def sample_parse_result(sample_parsed_resume, sample_metadata) -> ParseResult:
    return ParseResult(
        parsed_data=sample_parsed_resume,
        metadata=sample_metadata,
        status=ParseStatus.COMPLETED,
    )


@pytest.fixture
def sample_resume_text() -> str:
    return """\
Jane Doe
jane@example.com | +1-555-123-4567 | San Francisco, CA

SUMMARY
Experienced software engineer with 5+ years of experience in Python backends.

EXPERIENCE
Senior Engineer | Acme Corp | 2021 - Present
- Built microservices platform
- Reduced latency 40%

EDUCATION
MIT | BS Computer Science | 2015 - 2019 | GPA: 3.8

SKILLS
Python, FastAPI, PostgreSQL, React, Docker

PROJECTS
OpenAnalytics - Open-source analytics dashboard
https://github.com/janedoe/openanalytics
Technologies: Python, React

CERTIFICATIONS
AWS Solutions Architect - Amazon Web Services - 2023-03

LINKS
LinkedIn: https://linkedin.com/in/janedoe
GitHub: https://github.com/janedoe
"""


@pytest.fixture
def sample_docx_bytes() -> bytes:
    """Create a minimal DOCX file in memory."""
    doc = Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("jane@example.com | +1-555-123-4567 | San Francisco, CA")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Senior Engineer at Acme Corp (2021 - Present)")
    doc.add_paragraph("Built microservices platform")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("MIT | BS Computer Science | 2015 - 2019")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, FastAPI, PostgreSQL, React")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ── Mock Repository ─────────────────────────────────────────────────


@pytest.fixture
def mock_repository() -> MagicMock:
    repo = MagicMock(spec=ResumeRepository)
    repo.create = AsyncMock()
    repo.get_by_id = AsyncMock()
    repo.list_records = AsyncMock()
    repo.delete = AsyncMock()
    return repo


@pytest.fixture
def parser_service(mock_repository: MagicMock) -> ResumeParserService:
    return ResumeParserService(repository=mock_repository)
