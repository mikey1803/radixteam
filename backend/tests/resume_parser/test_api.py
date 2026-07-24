"""Tests for Resume Parser API endpoints."""

from __future__ import annotations

import io
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document

from shared.exceptions import ValidationException, NotFoundException

from backend.modules.resume_parser.api import router, set_parser_service
from backend.modules.resume_parser.schemas import (
    FileType,
    ParseMetadata,
    ParseResult,
    ParseStatus,
    ParsedResume,
)


def _make_result(**overrides) -> ParseResult:
    defaults = dict(
        parsed_data=ParsedResume(
            personal_info={
                "first_name": "Jane",
                "last_name": "Doe",
                "email": "jane@test.com",
            },
            summary="Engineer",
            skills=[{"name": "Python"}],
        ),
        metadata=ParseMetadata(
            parse_id="test-id",
            filename="resume.pdf",
            file_type=FileType.PDF,
            file_size_bytes=1024,
        ),
        status=ParseStatus.COMPLETED,
    )
    defaults.update(overrides)
    return ParseResult(**defaults)


def _make_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("jane@test.com | Engineer at Acme")
    doc.add_paragraph("Python, FastAPI, PostgreSQL")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


@pytest.fixture
def mock_service():
    svc = MagicMock()
    svc.parse_resume = AsyncMock()
    svc.get_parse_result = AsyncMock()
    svc.list_parse_history = AsyncMock()
    svc.delete_parse_result = AsyncMock()
    return svc


@pytest.fixture
def client(mock_service):
    app = FastAPI()
    app.include_router(router)
    set_parser_service(mock_service)
    return TestClient(app, raise_server_exceptions=False)


class TestParseEndpoint:
    def test_parses_docx(self, client, mock_service):
        mock_service.parse_resume.return_value = _make_result(
            filename="resume.docx",
            metadata=ParseMetadata(
                parse_id="new-id",
                filename="resume.docx",
                file_type=FileType.DOCX,
                file_size_bytes=2048,
            ),
        )
        docx_bytes = _make_docx_bytes()
        response = client.post(
            "/resume/parse",
            files={"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert response.status_code == 201
        body = response.json()
        assert body["success"] is True
        assert body["data"]["metadata"]["parse_id"] == "new-id"

    def test_validation_error(self, client, mock_service):
        mock_service.parse_resume.side_effect = ValidationException(
            message="Unsupported file type"
        )
        response = client.post(
            "/resume/parse",
            files={"file": ("resume.txt", b"content", "text/plain")},
        )
        assert response.status_code == 422
        assert response.json()["success"] is False


class TestGetParseResultEndpoint:
    def test_retrieves(self, client, mock_service):
        mock_service.get_parse_result.return_value = _make_result()
        response = client.get("/resume/parse/test-id")
        assert response.status_code == 200
        assert response.json()["data"]["metadata"]["parse_id"] == "test-id"

    def test_not_found(self, client, mock_service):
        mock_service.get_parse_result.side_effect = NotFoundException(
            "ParsedResumeRecord", "missing"
        )
        response = client.get("/resume/parse/missing")
        assert response.status_code == 404


class TestListParseHistoryEndpoint:
    def test_lists(self, client, mock_service):
        mock_service.list_parse_history.return_value = [_make_result(), _make_result()]
        response = client.get("/resume/parses")
        assert response.status_code == 200
        assert len(response.json()["data"]) == 2

    def test_empty_list(self, client, mock_service):
        mock_service.list_parse_history.return_value = []
        response = client.get("/resume/parses")
        assert response.status_code == 200
        assert response.json()["data"] == []


class TestDeleteParseResultEndpoint:
    def test_deletes(self, client, mock_service):
        mock_service.delete_parse_result.return_value = None
        response = client.delete("/resume/parse/test-id")
        assert response.status_code == 200
        assert response.json()["success"] is True

    def test_not_found(self, client, mock_service):
        mock_service.delete_parse_result.side_effect = NotFoundException(
            "ParsedResumeRecord", "missing"
        )
        response = client.delete("/resume/parse/missing")
        assert response.status_code == 404
