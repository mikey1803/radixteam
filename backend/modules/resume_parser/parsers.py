"""Low-level text extraction from PDF and DOCX files."""

from __future__ import annotations

import io
from pathlib import Path

from shared.logging import get_logger

logger = get_logger(__name__)

_MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from a resume file.

    Dispatches to the appropriate parser based on file extension.

    Raises:
        ValueError: If the file type is unsupported or extraction fails.
        ValueError: If the file is empty.
    """
    if not file_bytes:
        raise ValueError("File is empty")

    if len(file_bytes) > _MAX_FILE_SIZE_BYTES:
        raise ValueError(
            f"File size exceeds {_MAX_FILE_SIZE_BYTES // (1024 * 1024)} MB limit"
        )

    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext == ".pdf":
        return _extract_pdf(file_bytes)
    elif ext == ".docx":
        return _extract_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF file using pdfplumber."""
    import pdfplumber

    text_parts: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
    except Exception as exc:
        logger.error("PDF extraction failed", error=str(exc))
        raise ValueError(f"Failed to extract text from PDF: {exc}") from exc

    text = "\n".join(text_parts).strip()
    if not text:
        raise ValueError("No text content found in PDF")
    return text


def _extract_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file using python-docx."""
    from docx import Document

    text_parts: list[str] = []
    try:
        doc = Document(io.BytesIO(file_bytes))
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)
    except Exception as exc:
        logger.error("DOCX extraction failed", error=str(exc))
        raise ValueError(f"Failed to extract text from DOCX: {exc}") from exc

    text = "\n".join(text_parts).strip()
    if not text:
        raise ValueError("No text content found in DOCX")
    return text


def validate_file(filename: str, file_size: int) -> str:
    """Validate file type and size, return the detected extension."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext}'. Upload a PDF or DOCX file."
        )
    if file_size > _MAX_FILE_SIZE_BYTES:
        raise ValueError(
            f"File too large ({file_size} bytes). Maximum is {_MAX_FILE_SIZE_BYTES} bytes."
        )
    return ext
